#!/usr/bin/env python3
"""
Generate PDS Namespace Registry files (PDF and XLSX) from CSV.

This script reads pds-namespace-registry.csv and creates:
- pds-namespace-registry.xlsx
- pds-namespace-registry.pdf

Requirements:
    Python 3.9+
    pip install --quiet 'openpyxl~=3.1.5' 'reportlab~=4.5.1'
"""

from __future__ import annotations

import csv
import sys
from datetime import datetime
from pathlib import Path

# Third-party imports
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph
from reportlab.lib.enums import TA_LEFT


def read_csv_file(csv_path: Path) -> list[list[str]]:
    """Read the CSV file and return all rows."""
    rows = []
    with open(csv_path, 'r', encoding='utf-8-sig') as f:
        reader = csv.reader(f)
        for row in reader:
            rows.append(row)
    return rows


def create_xlsx(csv_rows: list[list[str]], output_path: Path) -> None:
    """Create an XLSX file with styling matching the actual file."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    # Define styles
    # Header row style (rows 1 and 2) - gray background, bold
    header_font = Font(name='Calibri', size=12, bold=True)
    header_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    header_alignment = Alignment(vertical='top', wrap_text=True)

    # Section header style (Common, International, Discipline, Mission) - light gray background, bold
    section_font = Font(name='Calibri', size=12, bold=True)
    section_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    section_alignment = Alignment(vertical='top', wrap_text=True)

    # Regular cell style
    regular_font = Font(name='Calibri', size=12)
    regular_alignment = Alignment(vertical='top', wrap_text=True)

    # Border style - thin borders on all sides
    thin_border = Border(
        left=Side(style='thin', color='000000'),
        right=Side(style='thin', color='000000'),
        top=Side(style='thin', color='000000'),
        bottom=Side(style='thin', color='000000')
    )

    # Write data
    for row_idx, csv_row in enumerate(csv_rows, start=1):
        # Ensure we have exactly 7 columns
        while len(csv_row) < 7:
            csv_row.append('')

        # Check if this row is a section header
        is_section_header = bool(csv_row) and csv_row[0].strip() in ['Common', 'International', 'Discipline', 'Mission', 'Held For Future Use']

        for col_idx, value in enumerate(csv_row[:7], start=1):
            cell = ws.cell(row=row_idx, column=col_idx)

            # Handle date values in row 1, column 2
            if row_idx == 1 and col_idx == 2 and value:
                try:
                    # Parse date from CSV (e.g., "4/15/2026")
                    date_obj = datetime.strptime(value.strip(), "%m/%d/%Y")
                    cell.value = date_obj
                    cell.number_format = 'M/D/YYYY'
                except (ValueError, AttributeError):
                    cell.value = value
            # Handle date values in column 7 (Registration Date)
            elif col_idx == 7 and value and row_idx > 2:
                try:
                    # Parse date from CSV (e.g., "2012-04-03")
                    date_obj = datetime.strptime(value.strip(), "%Y-%m-%d")
                    cell.value = date_obj
                    cell.number_format = 'YYYY-MM-DD'
                except (ValueError, AttributeError):
                    cell.value = value
            else:
                cell.value = value if value else None

            # Apply styling
            if row_idx <= 2:
                # Header rows (1 and 2)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
            elif is_section_header:
                # Section header row - apply to all columns A-G
                cell.font = section_font
                cell.fill = section_fill
                cell.alignment = section_alignment
            else:
                # Regular data cells
                cell.font = regular_font
                cell.alignment = regular_alignment

            # Apply borders to all cells
            cell.border = thin_border

    # Set column widths to match the actual file (only columns A-G)
    column_widths = {
        'A': 14.29,
        'B': 23.79,
        'C': 53.79,
        'D': 12.29,
        'E': 23.0,
        'F': 31.66,
        'G': 16.21
    }

    for col_letter, width in column_widths.items():
        ws.column_dimensions[col_letter].width = width

    # Save the workbook
    wb.save(output_path)
    print(f"Created XLSX file: {output_path}")


def create_pdf(csv_rows: list[list[str]], output_path: Path) -> None:
    """Create a PDF file with table layout matching the actual file."""
    doc = SimpleDocTemplate(
        str(output_path),
        pagesize=landscape(letter),
        leftMargin=0.5*inch,
        rightMargin=0.5*inch,
        topMargin=0.5*inch,
        bottomMargin=0.5*inch
    )

    # Container for the 'Flowable' objects
    elements = []

    # Define styles
    styles = getSampleStyleSheet()
    normal_style = ParagraphStyle(
        'Normal',
        parent=styles['Normal'],
        fontSize=8,
        leading=10,
        alignment=TA_LEFT
    )

    # Prepare table data with Paragraph objects for text wrapping
    table_data = []

    for row_idx, csv_row in enumerate(csv_rows):
        # Ensure we have exactly 7 columns
        while len(csv_row) < 7:
            csv_row.append('')

        # Convert text to Paragraph objects for better wrapping
        row_data = []
        for col_idx, cell_value in enumerate(csv_row[:7]):
            if cell_value:
                # Format dates in column 7
                if col_idx == 6 and row_idx > 1:  # Registration Date column
                    try:
                        date_obj = datetime.strptime(cell_value.strip(), "%Y-%m-%d")
                        cell_value = date_obj.strftime("%Y-%m-%d")
                    except (ValueError, AttributeError):
                        pass
                row_data.append(Paragraph(str(cell_value), normal_style))
            else:
                row_data.append('')

        table_data.append(row_data)

    # Create table
    # Column widths adjusted to fit on landscape letter (11" wide with 1" total margins = 10" available)
    col_widths = [0.9*inch, 1.4*inch, 2.8*inch, 0.8*inch, 1.3*inch, 1.6*inch, 0.9*inch]

    table = Table(table_data, colWidths=col_widths, repeatRows=2)

    # Define table style
    table_style = [
        # Header rows (rows 0 and 1)
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#D9D9D9')),
        ('BACKGROUND', (0, 1), (-1, 1), colors.HexColor('#D9D9D9')),
        ('TEXTCOLOR', (0, 0), (-1, 1), colors.black),
        ('FONTNAME', (0, 0), (-1, 1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 1), 9),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),

        # Grid
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ROWBACKGROUNDS', (0, 2), (-1, -1), [colors.white, colors.white]),

        # Default font for data rows
        ('FONTNAME', (0, 2), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 2), (-1, -1), 8),
    ]

    # Add section header styling (bold + light gray background)
    section_headers = ['Common', 'International', 'Discipline', 'Mission', 'Held For Future Use']
    for row_idx, csv_row in enumerate(csv_rows):
        if csv_row and csv_row[0].strip() in section_headers:
            table_style.append(('BACKGROUND', (0, row_idx), (-1, row_idx), colors.HexColor('#F2F2F2')))
            table_style.append(('FONTNAME', (0, row_idx), (-1, row_idx), 'Helvetica-Bold'))

    table.setStyle(TableStyle(table_style))
    elements.append(table)

    # Build PDF
    doc.build(elements)
    print(f"Created PDF file: {output_path}")


def main() -> int:
    """Main function to generate both XLSX and PDF files."""
    # Check Python version
    if sys.version_info < (3, 9):
        print("Error: Python 3.9 or higher is required")
        print(f"Current version: {sys.version}")
        return 1

    # Define paths
    repo_root = Path(__file__).parent.parent  # Go up from scripts/ to repo root
    namespace_registry_dir = repo_root / "docs" / "namespace-registry"
    csv_path = namespace_registry_dir / "pds-namespace-registry.csv"

    # Check if CSV exists
    if not csv_path.exists():
        print(f"Error: CSV file not found at {csv_path}")
        return 1

    # Read CSV
    print(f"Reading CSV file: {csv_path}")
    csv_rows = read_csv_file(csv_path)
    print(f"Read {len(csv_rows)} rows from CSV")

    # Generate XLSX
    xlsx_output = namespace_registry_dir / "pds-namespace-registry.xlsx"
    try:
        create_xlsx(csv_rows, xlsx_output)
    except Exception as e:
        print(f"Error creating XLSX: {e}")
        return 1

    # Generate PDF
    pdf_output = namespace_registry_dir / "pds-namespace-registry.pdf"
    try:
        create_pdf(csv_rows, pdf_output)
    except Exception as e:
        print(f"Error creating PDF: {e}")
        return 1

    print("\nSuccess! Generated files:")
    print(f"  - {xlsx_output}")
    print(f"  - {pdf_output}")
    return 0


if __name__ == "__main__":
    exit(main())
